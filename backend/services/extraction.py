"""
extraction.py
-------------
Why this module exists:
    Turns an uploaded file on disk into plain text, regardless of format.
    Kept separate from the upload router so the router only deals with
    HTTP concerns (size limits, saving the file) and this module only
    deals with "given a path, give me text" — easy to unit test and easy
    to extend with new formats later.

Supported formats: .txt, .md, .pdf, .docx, .pptx, .xlsx
    PDF/DOCX/PPTX/XLSX all ship real text layers in the vast majority of
    real-world files, so we extract that directly. Scanned (image-only)
    PDFs have no text layer — for those we fall back to OCR only if
    OCR_ENABLED=true in .env, since Tesseract is a separate system
    install we don't want to require by default.

Non-obvious design decisions:
    - Every extractor raises ExtractionError with a human-readable message
      instead of leaking library-specific exceptions. Callers (the upload
      router) can then store that message directly as the document's
      error detail — consistent with "never silently pretend it worked."
    - We never guess encoding beyond utf-8 with errors="replace" for plain
      text files; a slightly mangled character is fine, silently truncated
      text is not.
"""

from __future__ import annotations

from pathlib import Path

from backend.config import settings


class ExtractionError(Exception):
    """Raised when a file's text could not be extracted, with a message
    that is safe and useful to show directly to the end user."""


def extract_text(path: Path, original_filename: str) -> str:
    """
    Dispatch to the right extractor based on file extension.

    original_filename is used (not `path`) to decide the format, since
    uploaded files are saved on disk under a generated document id, not
    their original name.
    """
    suffix = Path(original_filename).suffix.lower()
    try:
        if suffix in (".txt", ".md"):
            text = _extract_plain_text(path)
        elif suffix == ".pdf":
            text = _extract_pdf(path)
        elif suffix == ".docx":
            text = _extract_docx(path)
        elif suffix == ".pptx":
            text = _extract_pptx(path)
        elif suffix == ".xlsx":
            text = _extract_xlsx(path)
        else:
            text = None
        if text is not None:
            return _strip_binary_junk(text)
    except ExtractionError:
        raise
    except Exception as exc:  # noqa: BLE001 - convert to our own error type
        raise ExtractionError(f"Failed to read {suffix} file: {exc}") from exc

    raise ExtractionError(
        f"Unsupported file type '{suffix}'. Supported: .txt, .md, .pdf, .docx, .pptx, .xlsx"
    )


def _strip_binary_junk(text: str) -> str:
    """Drop NUL / control bytes that break downstream tokenizers."""
    if not text:
        return ""
    text = text.replace("\x00", " ")
    text = "".join(ch if (ch >= " " or ch in "\n\t\r") else " " for ch in text)
    return text.encode("utf-8", errors="ignore").decode("utf-8")


def _extract_plain_text(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace")


def _extract_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages_text = [page.extract_text() or "" for page in reader.pages]
    text = "\n\n".join(pages_text).strip()

    if text:
        return text

    # No text layer found at all -> likely a scanned/image-only PDF.
    if settings.ocr_enabled:
        return _ocr_pdf(path)

    raise ExtractionError(
        "This PDF has no extractable text layer (it's likely a scanned "
        "image). Enable OCR_ENABLED=true in .env (requires Tesseract "
        "installed) to handle scanned PDFs."
    )


def _ocr_pdf(path: Path) -> str:
    """OCR fallback for image-only PDFs. Only reached when OCR_ENABLED=true."""
    try:
        import pytesseract
        from pdf2image import convert_from_path

        if settings.tesseract_cmd:
            pytesseract.pytesseract.tesseract_cmd = settings.tesseract_cmd

        images = convert_from_path(str(path))
        pages = [pytesseract.image_to_string(img) for img in images]
        text = "\n\n".join(pages).strip()
        if not text:
            raise ExtractionError("OCR ran but found no readable text in this PDF.")
        return text
    except ImportError as exc:
        raise ExtractionError(
            "OCR_ENABLED=true but pytesseract/pdf2image aren't installed. "
            "Run: pip install pytesseract pdf2image (and install the "
            "Tesseract OCR + poppler system packages)."
        ) from exc


def _extract_docx(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells)
            if row_text.strip(" |"):
                parts.append(row_text)
    text = "\n".join(parts).strip()
    if not text:
        raise ExtractionError("This .docx file has no readable text content.")
    return text


def _extract_pptx(path: Path) -> str:
    from pptx import Presentation

    presentation = Presentation(str(path))
    parts = []
    for i, slide in enumerate(presentation.slides, start=1):
        slide_lines = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                slide_lines.append(shape.text_frame.text.strip())
        if slide_lines:
            parts.append(f"[Slide {i}]\n" + "\n".join(slide_lines))
    text = "\n\n".join(parts).strip()
    if not text:
        raise ExtractionError("This .pptx file has no readable text content.")
    return text


def _extract_xlsx(path: Path) -> str:
    import openpyxl

    workbook = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    parts = []
    for sheet in workbook.worksheets:
        rows_text = []
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                rows_text.append(" | ".join(cells))
        if rows_text:
            parts.append(f"[Sheet: {sheet.title}]\n" + "\n".join(rows_text))
    text = "\n\n".join(parts).strip()
    if not text:
        raise ExtractionError("This .xlsx file has no readable cell content.")
    return text
